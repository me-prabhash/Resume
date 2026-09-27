"""
HTML to ATS-friendly PDF and DOCX converter.

Uses Playwright (headless Chromium) for PDF and python-docx for DOCX.
Parses the HTML resume to extract all sections and generates both formats.

Usage:
    python convert_to_pdf.py                        # generates both PDF and DOCX
    python convert_to_pdf.py --pdf-only              # PDF only
    python convert_to_pdf.py --docx-only             # DOCX only
    python convert_to_pdf.py path/to/resume.html     # custom input
"""

import sys
import os
import re
import logging
from html.parser import HTMLParser

logging.getLogger('pdfminer').setLevel(logging.ERROR)

# Defaults
DEFAULT_HTML = os.path.join(
    os.path.dirname(__file__), 'Resume', 'Prabhash_Thakur_Resume.html'
)
DEFAULT_PDF = os.path.join(
    os.path.dirname(__file__), 'Resume', 'Prabhash Thakur Resume.pdf'
)
DEFAULT_DOCX = os.path.join(
    os.path.dirname(__file__), 'Resume', 'Prabhash Thakur Resume.docx'
)


# ---------------------------------------------------------------------------
# HTML parser — extracts structured resume data from the HTML
# ---------------------------------------------------------------------------

class ResumeHTMLParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.data = {
            'name': '',
            'title': '',
            'contact': '',
            'summary': '',
            'skills': [],        # list of (label, text)
            'experience': [],    # list of {role, dates, company, bullets:[{text, bold_phrases:[]}]}
            'certifications': [],  # list of {name, issuer, url}
            'education': [],     # list of raw text strings
            'recognition': [],   # list of {text, bold_phrases:[]}
            'languages': '',
        }
        self._tag_stack = []
        self._class_stack = []
        self._current_text = ''
        self._current_section = None
        self._in_exp_item = False
        self._current_exp = None
        self._in_li = False
        self._li_text = ''
        self._li_bold_phrases = []
        self._in_strong = False
        self._strong_text = ''
        self._in_skills_row = False
        self._skills_label = ''
        self._skills_text = ''
        self._in_skills_label = False
        self._in_cert_item = False
        self._cert_name = ''
        self._cert_url = ''
        self._cert_rest = ''
        self._in_cert_link = False
        self._in_edu_item = False
        self._edu_text = ''
        self._in_award_item = False
        self._award_text = ''
        self._award_bold_phrases = []

    def handle_starttag(self, tag, attrs):
        attrs_dict = dict(attrs)
        cls = attrs_dict.get('class', '')
        self._tag_stack.append(tag)
        self._class_stack.append(cls)

        if cls == 'section-title':
            self._current_text = ''
        elif cls == 'name':
            self._current_text = ''
        elif tag == 'p' and cls == 'title':
            self._current_text = ''
        elif cls == 'contact-line':
            self._current_text = ''
        elif cls == 'summary-text':
            self._current_text = ''
        elif cls == 'skills-row':
            self._in_skills_row = True
            self._skills_label = ''
            self._skills_text = ''
        elif cls == 'skills-label':
            self._in_skills_label = True
            self._current_text = ''
        elif cls == 'experience-item':
            self._in_exp_item = True
            self._current_exp = {'role': '', 'dates': '', 'company': '', 'bullets': []}
        elif cls == 'exp-role':
            self._current_text = ''
        elif cls == 'exp-dates':
            self._current_text = ''
        elif cls == 'exp-company':
            self._current_text = ''
        elif tag == 'li' and self._in_exp_item:
            self._in_li = True
            self._li_text = ''
            self._li_bold_phrases = []
        elif tag == 'strong':
            self._in_strong = True
            self._strong_text = ''
        elif cls == 'cert-item':
            self._in_cert_item = True
            self._cert_name = ''
            self._cert_url = ''
            self._cert_rest = ''
        elif tag == 'a' and self._in_cert_item:
            self._in_cert_link = True
            self._cert_url = attrs_dict.get('href', '')
            self._current_text = ''
        elif cls == 'edu-item':
            self._in_edu_item = True
            self._edu_text = ''
        elif cls == 'award-item':
            self._in_award_item = True
            self._award_text = ''
            self._award_bold_phrases = []
        elif cls == 'lang-text':
            self._current_text = ''

    def handle_endtag(self, tag):
        if not self._tag_stack:
            return
        cls = self._class_stack[-1] if self._class_stack else ''

        if cls == 'section-title':
            title = self._current_text.strip().lower()
            if 'summary' in title:
                self._current_section = 'summary'
            elif 'skill' in title:
                self._current_section = 'skills'
            elif 'experience' in title:
                self._current_section = 'experience'
            elif 'certification' in title:
                self._current_section = 'certifications'
            elif 'education' in title:
                self._current_section = 'education'
            elif 'recognition' in title:
                self._current_section = 'recognition'
            elif 'language' in title:
                self._current_section = 'languages'
        elif cls == 'name':
            self.data['name'] = self._current_text.strip()
        elif tag == 'p' and cls == 'title':
            self.data['title'] = self._current_text.strip()
        elif cls == 'contact-line':
            self.data['contact'] = _normalize_whitespace(self._current_text)
        elif cls == 'summary-text':
            self.data['summary'] = _normalize_whitespace(self._current_text)
        elif cls == 'skills-label':
            self._in_skills_label = False
            self._skills_label = self._current_text.strip().rstrip(':')
        elif cls == 'skills-row':
            self._in_skills_row = False
            self.data['skills'].append((self._skills_label, _normalize_whitespace(self._skills_text)))
        elif cls == 'experience-item':
            self._in_exp_item = False
            if self._current_exp:
                self.data['experience'].append(self._current_exp)
            self._current_exp = None
        elif cls == 'exp-role':
            if self._current_exp is not None:
                self._current_exp['role'] = self._current_text.strip()
        elif cls == 'exp-dates':
            if self._current_exp is not None:
                self._current_exp['dates'] = self._current_text.strip()
        elif cls == 'exp-company':
            if self._current_exp is not None:
                self._current_exp['company'] = self._current_text.strip()
        elif tag == 'li' and self._in_li:
            self._in_li = False
            if self._current_exp is not None:
                self._current_exp['bullets'].append({
                    'text': _normalize_whitespace(self._li_text),
                    'bold_phrases': self._li_bold_phrases[:],
                })
        elif tag == 'strong':
            self._in_strong = False
            phrase = self._strong_text.strip()
            if phrase:
                if self._in_li:
                    self._li_bold_phrases.append(phrase)
                elif self._in_award_item:
                    self._award_bold_phrases.append(phrase)
        elif tag == 'a' and self._in_cert_link:
            self._in_cert_link = False
            self._cert_name = self._current_text.strip()
        elif cls == 'cert-item':
            self._in_cert_item = False
            issuer = self._cert_rest.strip().lstrip('-').lstrip(' \u2013').lstrip(' \u2014').strip()
            self.data['certifications'].append({
                'name': self._cert_name,
                'issuer': issuer,
                'url': self._cert_url,
            })
        elif cls == 'edu-item':
            self._in_edu_item = False
            self.data['education'].append(_normalize_whitespace(self._edu_text))
        elif cls == 'award-item':
            self._in_award_item = False
            self.data['recognition'].append({
                'text': _normalize_whitespace(self._award_text),
                'bold_phrases': self._award_bold_phrases[:],
            })
        elif cls == 'lang-text':
            self.data['languages'] = self._current_text.strip()

        self._tag_stack.pop()
        self._class_stack.pop()

    def handle_data(self, data):
        self._current_text += data
        if self._in_strong:
            self._strong_text += data
        if self._in_li:
            self._li_text += data
        if self._in_skills_row and not self._in_skills_label:
            self._skills_text += data
        if self._in_cert_item and not self._in_cert_link:
            self._cert_rest += data
        if self._in_edu_item:
            self._edu_text += data
        if self._in_award_item:
            self._award_text += data

    def handle_entityref(self, name):
        char_map = {'amp': '&', 'lt': '<', 'gt': '>', 'ndash': '\u2013', 'mdash': '\u2014', 'nbsp': ' '}
        ch = char_map.get(name, '')
        if ch:
            self.handle_data(ch)


def _normalize_whitespace(text):
    return re.sub(r'\s+', ' ', text).strip()


def parse_html_resume(html_path):
    with open(html_path, 'r', encoding='utf-8') as f:
        html = f.read()
    parser = ResumeHTMLParser()
    parser.feed(html)
    return parser.data


# ---------------------------------------------------------------------------
# PDF generation (Playwright)
# ---------------------------------------------------------------------------

def convert_html_to_pdf(html_path, pdf_path):
    from playwright.sync_api import sync_playwright

    html_path = os.path.abspath(html_path)
    pdf_path = os.path.abspath(pdf_path)

    if not os.path.exists(html_path):
        print(f"ERROR: HTML file not found: {html_path}")
        return False

    file_url = 'file:///' + html_path.replace('\\', '/')

    print(f"Converting to PDF: {html_path}")
    print(f"Output:            {pdf_path}")

    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        page.goto(file_url, wait_until='networkidle')
        page.wait_for_timeout(2000)
        page.pdf(
            path=pdf_path,
            format='A4',
            margin={
                'top': '10mm',
                'right': '12mm',
                'bottom': '10mm',
                'left': '12mm',
            },
            print_background=True,
        )
        browser.close()

    file_size = os.path.getsize(pdf_path)
    print(f"PDF created: {file_size:,} bytes")

    try:
        import pdfplumber
        with pdfplumber.open(pdf_path) as pdf:
            total_chars = sum(len(pg.extract_text() or '') for pg in pdf.pages)
            if total_chars > 100:
                print(f"Text check: PASS ({total_chars:,} characters extracted)")
            else:
                print(f"Text check: WARNING (only {total_chars} characters found)")
    except ImportError:
        print("(Install pdfplumber to verify text extraction)")

    return True


# ---------------------------------------------------------------------------
# DOCX generation (python-docx)
# ---------------------------------------------------------------------------

def convert_html_to_docx(html_path, docx_path):
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    data = parse_html_resume(html_path)

    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    ACCENT = RGBColor(0x16, 0x32, 0x5C)
    TEXT_DARK = RGBColor(0x11, 0x11, 0x11)
    TEXT_MEDIUM = RGBColor(0x2B, 0x2B, 0x2B)
    TEXT_LIGHT = RGBColor(0x5A, 0x5A, 0x5A)
    FONT_BODY = 'Calibri'
    FONT_HEADING = 'Calibri'

    def set_spacing(paragraph, before=0, after=0, line=None):
        fmt = paragraph.paragraph_format
        fmt.space_before = Pt(before)
        fmt.space_after = Pt(after)
        if line is not None:
            fmt.line_spacing = Pt(line)

    def add_separator(doc):
        p = doc.add_paragraph()
        set_spacing(p, before=4, after=4)
        run = p.add_run('\u2500' * 80)
        run.font.size = Pt(6)
        run.font.color.rgb = RGBColor(0xCF, 0xCF, 0xCF)

    def add_section_heading(doc, text):
        add_separator(doc)
        p = doc.add_paragraph()
        set_spacing(p, before=2, after=4)
        run = p.add_run(text.upper())
        run.font.name = FONT_HEADING
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = ACCENT

    def add_bullet(doc, text, bold_phrases=None):
        p = doc.add_paragraph(style='List Bullet')
        set_spacing(p, before=0, after=1, line=14)
        if bold_phrases:
            remaining = text
            for phrase in bold_phrases:
                idx = remaining.find(phrase)
                if idx == -1:
                    continue
                if idx > 0:
                    run = p.add_run(remaining[:idx])
                    run.font.name = FONT_BODY
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = TEXT_MEDIUM
                run = p.add_run(phrase)
                run.font.name = FONT_BODY
                run.font.size = Pt(9.5)
                run.font.color.rgb = TEXT_MEDIUM
                run.bold = True
                remaining = remaining[idx + len(phrase):]
            if remaining:
                run = p.add_run(remaining)
                run.font.name = FONT_BODY
                run.font.size = Pt(9.5)
                run.font.color.rgb = TEXT_MEDIUM
        else:
            run = p.add_run(text)
            run.font.name = FONT_BODY
            run.font.size = Pt(9.5)
            run.font.color.rgb = TEXT_MEDIUM
        return p

    # --- HEADER ---
    p = doc.add_paragraph()
    set_spacing(p, before=0, after=0)
    run = p.add_run(data['name'])
    run.font.name = FONT_HEADING
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = TEXT_DARK

    p = doc.add_paragraph()
    set_spacing(p, before=0, after=2)
    run = p.add_run(data['title'])
    run.font.name = FONT_BODY
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = ACCENT

    p = doc.add_paragraph()
    set_spacing(p, before=0, after=2)
    run = p.add_run(data['contact'])
    run.font.name = FONT_BODY
    run.font.size = Pt(9)
    run.font.color.rgb = TEXT_MEDIUM

    # --- SUMMARY ---
    add_section_heading(doc, 'Professional Summary')
    p = doc.add_paragraph()
    set_spacing(p, before=0, after=2, line=14)
    run = p.add_run(data['summary'])
    run.font.name = FONT_BODY
    run.font.size = Pt(9.5)
    run.font.color.rgb = TEXT_MEDIUM

    # --- SKILLS ---
    add_section_heading(doc, 'Skills')
    for label, text in data['skills']:
        p = doc.add_paragraph()
        set_spacing(p, before=0, after=2, line=14)
        run = p.add_run(f"{label}: ")
        run.font.name = FONT_BODY
        run.font.size = Pt(9.5)
        run.font.color.rgb = TEXT_DARK
        run.bold = True
        run = p.add_run(text)
        run.font.name = FONT_BODY
        run.font.size = Pt(9.5)
        run.font.color.rgb = TEXT_MEDIUM

    # --- EXPERIENCE ---
    add_section_heading(doc, 'Work Experience')
    for exp in data['experience']:
        p = doc.add_paragraph()
        set_spacing(p, before=6, after=0)
        run = p.add_run(f"{exp['role']}  |  {exp['dates']}")
        run.font.name = FONT_BODY
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = TEXT_DARK

        p = doc.add_paragraph()
        set_spacing(p, before=0, after=2)
        run = p.add_run(exp['company'])
        run.font.name = FONT_BODY
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = ACCENT

        for bullet in exp['bullets']:
            add_bullet(doc, bullet['text'], bullet.get('bold_phrases'))

    # --- CERTIFICATIONS ---
    add_section_heading(doc, 'Certifications')
    for cert in data['certifications']:
        p = doc.add_paragraph(style='List Bullet')
        set_spacing(p, before=0, after=1, line=14)
        run = p.add_run(cert['name'])
        run.font.name = FONT_BODY
        run.font.size = Pt(9.5)
        run.font.color.rgb = TEXT_DARK
        run.bold = True
        if cert['issuer']:
            run = p.add_run(f" - {cert['issuer']}")
            run.font.name = FONT_BODY
            run.font.size = Pt(9.5)
            run.font.color.rgb = TEXT_MEDIUM

    # --- EDUCATION ---
    add_section_heading(doc, 'Education')
    for edu_text in data['education']:
        p = doc.add_paragraph()
        set_spacing(p, before=0, after=2, line=14)
        # Split on the ndash to bold the degree part
        parts = edu_text.split('\u2013', 1)
        if len(parts) == 2:
            run = p.add_run(parts[0].strip())
            run.font.name = FONT_BODY
            run.font.size = Pt(9.5)
            run.font.color.rgb = TEXT_DARK
            run.bold = True
            run = p.add_run(f" \u2013 {parts[1].strip()}")
            run.font.name = FONT_BODY
            run.font.size = Pt(9.5)
            run.font.color.rgb = TEXT_MEDIUM
        else:
            run = p.add_run(edu_text)
            run.font.name = FONT_BODY
            run.font.size = Pt(9.5)
            run.font.color.rgb = TEXT_MEDIUM

    # --- RECOGNITION ---
    add_section_heading(doc, 'Recognition')
    for award in data['recognition']:
        add_bullet(doc, award['text'], award.get('bold_phrases'))

    # --- LANGUAGES ---
    add_section_heading(doc, 'Languages')
    p = doc.add_paragraph()
    set_spacing(p, before=0, after=0)
    run = p.add_run(data['languages'])
    run.font.name = FONT_BODY
    run.font.size = Pt(9.5)
    run.font.color.rgb = TEXT_MEDIUM

    doc.save(docx_path)
    file_size = os.path.getsize(docx_path)
    print(f"DOCX created: {file_size:,} bytes  ->  {docx_path}")
    return True


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

if __name__ == '__main__':
    args = sys.argv[1:]
    pdf_only = '--pdf-only' in args
    docx_only = '--docx-only' in args
    args = [a for a in args if not a.startswith('--')]

    html_path = args[0] if args else DEFAULT_HTML
    pdf_path = args[1] if len(args) > 1 else DEFAULT_PDF
    docx_path = args[2] if len(args) > 2 else DEFAULT_DOCX

    success = True

    if not docx_only:
        print("=== PDF ===")
        success = convert_html_to_pdf(html_path, pdf_path) and success
        print()

    if not pdf_only:
        print("=== DOCX ===")
        success = convert_html_to_docx(html_path, docx_path) and success
        print()

    if success:
        print("Done!")
    else:
        sys.exit(1)
