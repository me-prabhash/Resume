"""
HTML Resume to Word (.docx) converter.

Parses the two-column ATS-friendly HTML resume and creates a clean,
single-column ATS-friendly Word document with proper headings, bullet points,
and formatting.

Linearizes the two-column layout into logical reading order:
Header → Summary → Skills → Experience → Certifications → Education → Languages → Recognition

Usage:
    python convert_to_docx.py
    python convert_to_docx.py path/to/resume.html path/to/output.docx
"""

import sys
import os
import re
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

DEFAULT_HTML = os.path.join(
    os.path.dirname(__file__), 'Resume', 'Prabhash_Thakur_Resume.html'
)
DEFAULT_DOCX = os.path.join(
    os.path.dirname(__file__), 'Resume', 'Prabhash Thakur Resume.docx'
)

FONT_NAME = 'Calibri'
COLOR_DARK = RGBColor(0x1A, 0x1A, 0x1A)
COLOR_MEDIUM = RGBColor(0x44, 0x44, 0x44)
COLOR_LIGHT = RGBColor(0x77, 0x77, 0x77)
COLOR_ACCENT = RGBColor(0x21, 0x96, 0xF3)
COLOR_HEADING = RGBColor(0x2D, 0x2D, 0x2D)


def clean_text(text):
    """Collapse all whitespace/newlines into a single space, like a browser does."""
    if not text:
        return ""
    return re.sub(r'\s+', ' ', text).strip()


def set_run(run, size=9.5, bold=False, color=COLOR_DARK, font=FONT_NAME):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_heading_text(doc, text, size=10, color=COLOR_HEADING):
    p = doc.add_paragraph()
    p.space_before = Pt(8)
    p.space_after = Pt(4)
    run = p.add_run(text.upper())
    set_run(run, size=size, bold=True, color=color)
    return p


def add_separator(doc):
    p = doc.add_paragraph()
    p.space_before = Pt(2)
    p.space_after = Pt(2)
    run = p.add_run('─' * 80)
    set_run(run, size=7, color=RGBColor(0x88, 0x88, 0x88))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_bullet(doc, text):
    """Handle <strong> tags by splitting into bold/normal runs."""
    # Collapse HTML newlines/spaces before regex splitting
    text = clean_text(text)
    parts = re.split(r'(<strong>.*?</strong>)', text)
    p = doc.add_paragraph(style='List Bullet')
    p.space_before = Pt(0)
    p.space_after = Pt(1)
    
    # Fix indentation so wrapped lines align with bullet text, not indent further
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    
    for part in parts:
        if part.startswith('<strong>'):
            clean = part.replace('<strong>', '').replace('</strong>', '')
            if clean:
                run = p.add_run(clean)
                set_run(run, size=9, bold=True, color=COLOR_MEDIUM)
        elif part.strip():
            clean = re.sub(r'<[^>]+>', '', part)
            if clean:
                run = p.add_run(clean)
                set_run(run, size=9, color=COLOR_MEDIUM)


def process_header(doc, soup):
    """Extract and render the header: Name, Title, Contact."""
    # Name
    name_el = soup.find('h1', class_='name')
    name = clean_text(name_el.get_text()) if name_el else 'PRABHASH THAKUR'
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(name)
    set_run(run, size=22, bold=True, color=COLOR_DARK)

    # Title
    title_el = soup.find('p', class_='title')
    title = clean_text(title_el.get_text()) if title_el else ''
    if title:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.space_before = Pt(0)
        run = p.add_run(title)
        set_run(run, size=10, color=COLOR_ACCENT)

    # Contact line
    contact_el = soup.find(class_='contact-line')
    if contact_el:
        contact_items = contact_el.find_all('span', class_='contact-item')
        if contact_items:
            parts = []
            for item in contact_items:
                text = clean_text(item.get_text())
                text = re.sub(r'[📞✉🔗📅]', '', text).strip()
                if text:
                    parts.append(text)
            contact_text = '  |  '.join(parts)
        else:
            contact_text = clean_text(contact_el.get_text())

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.space_before = Pt(2)
        run = p.add_run(contact_text)
        set_run(run, size=9, color=COLOR_MEDIUM)

    add_separator(doc)


def process_summary(doc, soup):
    """Extract and render the Summary/Profile section."""
    summary = soup.find('p', class_='summary-text') or soup.find('p', class_='profile-text')
    if summary:
        add_heading_text(doc, 'Professional Summary')
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(0)
        p.paragraph_format.first_line_indent = Pt(0)
        run = p.add_run(clean_text(summary.get_text()))
        set_run(run, size=9.5, color=COLOR_MEDIUM)
        add_separator(doc)


def process_skills(doc, soup):
    """Extract and render Skills section."""
    skill_groups = soup.find_all('div', class_='skills-category')
    if not skill_groups:
        skill_groups = soup.find_all('div', class_='skills-group')
    if not skill_groups:
        return

    add_heading_text(doc, 'Skills')

    for sg in skill_groups:
        title_div = sg.find('div', class_='skills-cat-title') or sg.find('div', class_='skills-group-title')
        tags = sg.find_all('span', class_='skill-tag')
        items_div = sg.find('div', class_='skills-group-items')

        if title_div:
            p = doc.add_paragraph()
            p.space_before = Pt(2)
            p.space_after = Pt(1)
            run = p.add_run(clean_text(title_div.get_text()) + ': ')
            set_run(run, size=9, bold=True, color=COLOR_DARK)

            if tags:
                skill_text = ', '.join(clean_text(tag.get_text()) for tag in tags)
            elif items_div:
                skill_text = clean_text(items_div.get_text())
            else:
                skill_text = ''

            if skill_text:
                run = p.add_run(skill_text)
                set_run(run, size=9, color=COLOR_MEDIUM)

    add_separator(doc)


def process_experience(doc, soup):
    """Extract and render Work Experience section."""
    exp_items = soup.find_all('div', class_='experience-item')
    if not exp_items:
        return

    add_heading_text(doc, 'Work Experience')

    for i, exp in enumerate(exp_items):
        if i > 0:
            p = doc.add_paragraph()
            p.space_before = Pt(4)

        role_el = exp.find('div', class_='exp-role')
        dates_el = exp.find('div', class_='exp-dates')
        company_el = exp.find('div', class_='exp-company')

        # Role + Dates on same line
        p = doc.add_paragraph()
        p.space_before = Pt(6)
        p.space_after = Pt(0)
        role = clean_text(role_el.get_text()) if role_el else ''
        dates = clean_text(dates_el.get_text()) if dates_el else ''
        dates = re.sub(r'[📅]', '', dates).strip()

        run = p.add_run(role)
        set_run(run, size=11, bold=True, color=COLOR_DARK)
        if dates:
            run = p.add_run(f'  |  {dates}')
            set_run(run, size=9, color=COLOR_LIGHT)

        # Company
        if company_el:
            p = doc.add_paragraph()
            p.space_before = Pt(0)
            p.space_after = Pt(2)
            run = p.add_run(clean_text(company_el.get_text()))
            set_run(run, size=9.5, color=COLOR_ACCENT)

        # Bullets
        bullets = exp.find_all('li')
        for li in bullets:
            add_bullet(doc, str(li.decode_contents()))

    add_separator(doc)


def process_certifications(doc, soup):
    """Extract and render Certifications section."""
    cert_items = soup.find_all('div', class_='cert-item')
    cert_list = soup.find('ul', class_='cert-list')

    if cert_items:
        add_heading_text(doc, 'Certifications')
        for ci in cert_items:
            name_el = ci.find('div', class_='cert-name')
            issuer_el = ci.find('div', class_='cert-issuer')
            if name_el:
                link = name_el.find('a')
                text = clean_text(link.get_text()) if link else clean_text(name_el.get_text())
                p = doc.add_paragraph(style='List Bullet')
                p.space_before = Pt(0)
                p.space_after = Pt(1)
                run = p.add_run(text)
                set_run(run, size=9, bold=True, color=COLOR_DARK)
                if issuer_el:
                    run = p.add_run(f' — {clean_text(issuer_el.get_text())}')
                    set_run(run, size=9, color=COLOR_LIGHT)
        add_separator(doc)
    elif cert_list:
        add_heading_text(doc, 'Certifications')
        for li in cert_list.find_all('li'):
            link = li.find('a')
            text = clean_text(link.get_text()) if link else clean_text(li.get_text())
            p = doc.add_paragraph(style='List Bullet')
            p.space_before = Pt(0)
            p.space_after = Pt(1)
            run = p.add_run(text)
            set_run(run, size=9, color=COLOR_DARK)
        add_separator(doc)


def process_education(doc, soup):
    """Extract and render Education section."""
    edu_items = soup.find_all('div', class_='edu-item')
    if not edu_items:
        return

    add_heading_text(doc, 'Education')

    for edu in edu_items:
        degree_el = edu.find('div', class_='edu-degree')
        school_el = edu.find('div', class_='edu-school')
        if degree_el:
            p = doc.add_paragraph()
            p.space_before = Pt(2)
            p.space_after = Pt(0)
            run = p.add_run(clean_text(degree_el.get_text()))
            set_run(run, size=9.5, bold=True, color=COLOR_DARK)
        if school_el:
            p = doc.add_paragraph()
            p.space_before = Pt(0)
            p.space_after = Pt(1)
            run = p.add_run(clean_text(school_el.get_text()))
            set_run(run, size=9, color=COLOR_LIGHT)

    add_separator(doc)


def process_languages(doc, soup):
    """Extract and render Languages section."""
    lang_items = soup.find_all('div', class_='lang-item')
    if lang_items:
        add_heading_text(doc, 'Languages')
        langs = []
        for li in lang_items:
            name_el = li.find('div', class_='lang-name')
            level_el = li.find('div', class_='lang-level')
            if name_el:
                name = clean_text(name_el.get_text())
                level = clean_text(level_el.get_text()) if level_el else ''
                langs.append(f'{name} ({level})' if level else name)
        if langs:
            p = doc.add_paragraph()
            run = p.add_run(', '.join(langs))
            set_run(run, size=9, color=COLOR_MEDIUM)
        add_separator(doc)
        return

    # Fall back: look for a Languages section with a plain <p>
    for sec in soup.find_all('section', class_='section'):
        title_el = sec.find('h2', class_='section-title')
        if title_el and 'language' in clean_text(title_el.get_text()).lower():
            lang_p = sec.find('p')
            if lang_p:
                add_heading_text(doc, 'Languages')
                p = doc.add_paragraph()
                run = p.add_run(clean_text(lang_p.get_text()))
                set_run(run, size=9, color=COLOR_MEDIUM)
                add_separator(doc)
            break


def process_recognition(doc, soup):
    """Extract and render Recognition section."""
    awards = soup.find_all('div', class_='award-item')
    if not awards:
        awards = soup.find_all('div', class_='achievement-item')

    if awards:
        add_heading_text(doc, 'Recognition')
        for ach in awards:
            add_bullet(doc, str(ach.decode_contents()))


def convert_html_to_docx(html_path, docx_path):
    html_path = os.path.abspath(html_path)
    docx_path = os.path.abspath(docx_path)

    if not os.path.exists(html_path):
        print(f"ERROR: HTML file not found: {html_path}")
        return False

    with open(html_path, 'r', encoding='utf-8') as f:
        soup = BeautifulSoup(f.read(), 'html.parser')

    doc = Document()

    # Set default font
    style = doc.styles['Normal']
    style.font.name = FONT_NAME
    style.font.size = Pt(9.5)
    style.paragraph_format.space_after = Pt(2)

    # Narrow margins for more content space
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    # Process sections in logical order for ATS readability
    process_header(doc, soup)
    process_summary(doc, soup)
    process_skills(doc, soup)
    process_experience(doc, soup)
    process_certifications(doc, soup)
    process_education(doc, soup)
    process_languages(doc, soup)
    process_recognition(doc, soup)

    doc.save(docx_path)
    file_size = os.path.getsize(docx_path)
    print(f"Converting: {html_path}")
    print(f"Output:     {docx_path}")
    print(f"Word doc created: {file_size:,} bytes")
    print("This .docx is ATS-readable and ready to send to HR.")
    return True


if __name__ == '__main__':
    html_path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_HTML
    docx_path = sys.argv[2] if len(sys.argv) > 2 else DEFAULT_DOCX

    success = convert_html_to_docx(html_path, docx_path)
    if not success:
        sys.exit(1)
